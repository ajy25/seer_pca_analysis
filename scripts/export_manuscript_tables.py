import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

PROJECT_ROOT = Path(__file__).resolve().parents[1]
TABLE_DIR = PROJECT_ROOT / "output" / "tables"
WORKBOOK_PATH = PROJECT_ROOT / "output" / "all_manuscript_tables.xlsx"
TABLE_FILES = {
    "Table 1 Characteristics": "12_patient_characteristics.csv",
    "Table 2 Survival by Era": "13_median_survival_by_era_race.csv",
    "Table 3 OS Cox": "14_cox_os.csv",
    "Supp Table 1 PH Config": "22_supplementary_table_1_ph_configurations.csv",
    "Supp Table 2 Time Varying": "22_supplementary_table_2_time_varying_coefficients.csv",
    "Supp Table 3 PCSS Cox": "15_cox_pcss.csv",
    "Supp Table 4 Cox by Stage": "16_cox_by_stage.csv",
    "Supp Table 5 Survival Stage": "18_median_survival_by_stage_race_era.csv",
    "Supp Table 6 Cox by Race": "17_cox_by_race.csv",
    "Supp Table 7 Era Race Int": "22_era_race_interaction.csv",
    "Supp Table 8 Cox Stage Race": "19_cox_by_stage_race.csv",
    "Supp Table 9 Sensitivity": "24_supplementary_table_9.csv",
}


def export_tables(table_dir=TABLE_DIR, output_path=WORKBOOK_PATH):
    missing = [
        name for name in TABLE_FILES.values() if not (table_dir / name).is_file()
    ]
    if missing:
        raise FileNotFoundError(
            "Generate the missing tables before exporting: " + ", ".join(missing)
        )
    workbook = Workbook()
    workbook.remove(workbook.active)
    for sheet_name, filename in TABLE_FILES.items():
        sheet = workbook.create_sheet(sheet_name)
        with (table_dir / filename).open(newline="", encoding="utf-8") as source:
            rows = list(csv.reader(source))
        if not rows or not rows[0] or any(len(row) != len(rows[0]) for row in rows):
            raise ValueError(f"Expected a nonempty rectangular table: {filename}")
        for row_index, row in enumerate(rows, start=1):
            for column_index, value in enumerate(row, start=1):
                cell = sheet.cell(row_index, column_index, value)
                cell.data_type = "s"
                if row_index == 1:
                    cell.font = Font(bold=True)
        sheet.freeze_panes = "B2"
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    print(f"Exported {len(TABLE_FILES)} tables to {output_path}", flush=True)


def main():
    export_tables()


if __name__ == "__main__":
    main()
